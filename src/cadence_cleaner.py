import pandas as pd
from pathlib import Path
import datetime
import openpyxl
import numpy as np
import re
from openpyxl.worksheet.table import Table, TableStyleInfo

# =========================================================
# CONFIGURATION & CONSTANTS
# =========================================================

BASE_FOLDER = Path(__file__).resolve().parent.parent
OUTPUT_DIR = Path.home() / "Desktop" / "Cleaned Data"

REQUIRED_COLUMNS = [
    "STORE_NAME",
    "BRAND",
    "EXT_PRICE_W_TAX",
    "RECEIPT_NO",
    "CUSTOMER_NAME",
]

COLUMNS_TO_DROP = [
    "SERIAL_NO",
    "MOBILE_NUMBER",
    "DATABASE_NAME",
    "CARD_TYPE",
    "TYPE_OF_PAYMENT",
    "TERMS",
    "AUTHORIZATION_CODE",
    "USE_CASE",
    "UNIT_ORIG_PRICE_W_TAX",
    "UNIT_DISC_AMT_W_TAX",
    "UNIT_FINAL_PRICE_W_TAX",
]

STORE_TYPO_FIXES = {
    "ALTITUDE PIONER": "ALTITUDE PIONEER",
}

DIRECT_STORES = [
    "ECOFLOW",
    "OURA",
    "RAPTOR",
    "INSTA360",
]

EF_TAG_DEFAULT = "EF Others"
I360_TAG_DEFAULT = "I360 Others"

EF_EXCLUDE_TERMS = [
    "CABLE", "CABLES", "ADAPTER", "ADAPTOR", "CHARGER", "FLASHLIGHT", 
    "BAG", "COVER", "TROLLEY", "PROTECTOR", "MOUNT", "DC"
]

BRAND_RULES = {
    "ALTITUDE PIONEER": ["DJI", "INSTA360", "ECOFLOW", "RAPTOR", "YABER", "ENABOT"],
    "ZENITH": ["RINGCONN","OURA","BOOX","COROS","PLAUDAI","GARMIN","SHOKZ","LOOP","POLAR","RAYNEO","EARQUIS"],
    "ALTITUDE": ["ECOFLOW", "INSTA360", "DJI","RINGCONN", "OURA", "BOOX", "COROS","PLAUDAI", "RAPTOR","SHOKZ","YABER","FUJIFILM", "RAYNEO","ENABOT","POLAR","GARMIN","OCCULI","EARQUIS","VULCROX","SONY"],
    "UG": ["ECOFLOW", "INSTA360", "DJI", "RINGCONN","ANKER","OURA","SHOKZ","YABER","GARMIN","PLAUDAI","RAYNEO","BOOX","COROS","POLAR","RAPTOR","OCCULLI","EARQUIS","SONY"],
}


# =========================================================
# HELPER FUNCTIONS: TEXT & HEADER PROCESSING
# =========================================================

def normalize_header(value) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip().upper().replace("\n", " ")


def clean_text(series: pd.Series) -> pd.Series:
    return (
        series.fillna("")
        .astype(str)
        .str.strip()
        .str.upper()
        .str.replace(r"\s+", " ", regex=True)
    )


def contains_term(text: str, term: str) -> bool:
    pattern = rf"(?<![A-Z0-9]){re.escape(term)}(?![A-Z0-9])"
    return bool(re.search(pattern, text))


def fix_store_typo(store: str) -> str:
    for wrong, correct in STORE_TYPO_FIXES.items():
        if store == wrong or store.startswith(wrong + " "):
            return correct + store[len(wrong):]
    return store


def read_data_file(file_path: Path, header=None, nrows=None) -> pd.DataFrame:
    if file_path.suffix.lower() in [".xls", ".xlsx"]:
        return pd.read_excel(file_path, header=header, nrows=nrows)
    return pd.read_csv(file_path, header=header, nrows=nrows)


# =========================================================
# MONOBRAND TAGGING LOGIC
# =========================================================

EF_TAG_DEFAULT = "EF Others"
I360_TAG_DEFAULT = "I360 Others"
RAPTOR_TAG_DEFAULT = "RAPTOR OTHERS"
OURA_TAG_DEFAULT = "OURA OTHERS"


def create_raptor_tag(long_description: str) -> str:
    """Classify Raptor products using LONG_DESCRIPTION keyword priority."""
    desc = str(long_description).upper()

    if any(k in desc for k in ["TOOLBOX", "TOOL BOX"]):
        return "TOOLBOX"
    elif "COOLER" in desc:
        return "COOLERS"
    elif any(k in desc for k in ["DRINKWARE", "BOTTLE", "MUG", "TUMBLER", "FLASK", "CUP"]):
        return "DRINKWARE"
    elif any(k in desc for k in ["APPAREL", "SHIRT", "HOODIE", "HAT", "CAP", "TEE", "JACKET", "POLO", "SHORTS", "PANTS"]):
        return "APPAREL"
    elif any(k in desc for k in ["FOODWARE", "LUNCH", "CUTLERY", "PLATE", "BOWL", "SPORK", "MESS KIT"]):
        return "FOODWARE"
    elif any(k in desc for k in ["CAMPING", "OUTING", "TENT", "CHAIR", "TABLE", "LANTERN", "HAMMOCK"]):
        return "OUTING"
    # Hardware takes priority over BAGS because "CARRY" appears in both rule sets.
    elif any(k in desc for k in ["CASE", "LITE", "EXTREME", "CARRY"]):
        return "HARDWARE"
    elif any(k in desc for k in ["BAG", "BACKPACK", "POUCH", "SLING", "DUFFLE", "PACK", "CARRY", "SACK"]):
        return "BAGS"
    elif any(k in desc for k in ["STORAGE", "BOX", "CONTAINER", "BIN"]):
        return "STORAGES"
    return RAPTOR_TAG_DEFAULT


def create_ef_tag(long_description: str) -> str:
    """Classify EcoFlow products; accessories/unmatched products go to EF Others."""
    desc = str(long_description).upper()

    # Accessories should not be classified as a power station series just because
    # their descriptions also mention a compatible River/Delta product.
    accessory_terms = ["CABLE", "CABLES", "ADAPTOR", "ADAPTORS", "ADAPTER", "ADAPTERS", "FLASHLIGHT", "FLASHLIGHTS"]
    if any(term in desc for term in accessory_terms):
        return EF_TAG_DEFAULT
    if "SOLAR" in desc and any(term in desc for term in ["PANEL", "PANELS", "SOLAR"]):
        return "Solar Panel"
    elif "DELTA" in desc:
        return "Delta Series"
    elif "RIVER" in desc:
        return "River Series"
    return EF_TAG_DEFAULT


def create_i360_tag(long_description: str) -> str:
    """Only tag a recognized Insta360 model when the description is a bundle or limited edition."""
    desc = str(long_description).upper()
    if "BUNDLE" not in desc and "LIMITED EDITION" not in desc:
        return I360_TAG_DEFAULT

    # Check more specific/multi-word model names before shorter model tokens.
    if "ACE PRO 2" in desc:
        return "Ace Pro 2"
    elif "GO ULTRA" in desc:
        return "Go Ultra"
    elif "GO 3S" in desc:
        return "Go 3s"
    elif "LUNA" in desc:
        return "Luna"
    elif "X5" in desc:
        return "X5"
    elif "X6" in desc:
        return "X6"
    return I360_TAG_DEFAULT


def create_oura_tag(long_description: str) -> str:
    """Classify Oura Ring 4, Ring 5, and accessory products."""
    desc = str(long_description).upper()
    if any(k in desc for k in ["CHARGER", "PROTECTOR", "SIZING KIT", "SIZING", "ACCESSORY", "ACCESSORIES"]):
        return "OURA ACCESSORIES"
    elif any(k in desc for k in ["RING 4", "RING4"]):
        return "OURA RING 4"
    elif any(k in desc for k in ["RING 5", "RING5"]):
        return "OURA RING 5"
    return OURA_TAG_DEFAULT


def resolve_product_tag(brand: str, long_description: str) -> str:
    """Choose tagging rules strictly from the BRAND cell, not store/category."""
    bnd = str(brand).strip().upper()
    if bnd == "RAPTOR":
        return create_raptor_tag(long_description)
    elif bnd == "ECOFLOW":
        return create_ef_tag(long_description)
    elif bnd == "INSTA360":
        return create_i360_tag(long_description)
    elif bnd == "OURA":
        return create_oura_tag(long_description)
    return ""


# =========================================================
# HELPER FUNCTIONS: DATA CLEANING & FILTERING
# =========================================================

def create_channel(store: str) -> str:
    if not store:
        return ""
    if store.startswith("ALTITUDE PIONEER"):
        return "ALTITUDE PIONEER"
    return store.split()[0]


def create_category(store: str, brand: str) -> str:
    if not store:
        return ""
    for name in DIRECT_STORES:
        if contains_term(store, name):
            return name
    for store_key, allowed_brands in BRAND_RULES.items():
        if contains_term(store, store_key):
            if not brand:
                return "OTHER BRANDS"
            for item in allowed_brands:
                if contains_term(brand, item):
                    return item

            return "OTHER BRANDS"
    return ""


def clean_special_brands(df: pd.DataFrame) -> pd.DataFrame:
    if "BRAND" not in df.columns or "LONG_DESCRIPTION" not in df.columns:
        return df
    brand = df["BRAND"].fillna("").astype(str).str.strip()
    long_desc = df["LONG_DESCRIPTION"].fillna("").astype(str)

    # Reclassify JINKO entries as RAPTOR when their long description mentions Raptor.
    jinko_raptor_mask = (
        brand.str.upper().eq("JINKO")
        & long_desc.str.contains("Raptor", case=False, na=False)
    )
    df.loc[jinko_raptor_mask, "BRAND"] = "RAPTOR"

    df.loc[(brand == "*") & long_desc.str.contains("Akaso", case=False, na=False), "BRAND"] = "Akaso"
    df.loc[brand.str.upper() == "PLAUD", "BRAND"] = "PlaudAi"
    
    brand = df["BRAND"].fillna("").astype(str).str.strip()
    mask_zz = brand == "ZZ"
    if mask_zz.any():
        has_paper_bag = long_desc.str.contains("Paper Bag", case=False, na=False)
        if "EXT_PRICE_W_TAX" in df.columns:
            zero_price = df["EXT_PRICE_W_TAX"].fillna(0) == 0
            df = df[~(mask_zz & has_paper_bag & zero_price)].reset_index(drop=True)
            brand = df["BRAND"].fillna("").astype(str).str.strip()
            long_desc = df["LONG_DESCRIPTION"].fillna("").astype(str)
            mask_zz = brand == "ZZ"
            has_paper_bag = long_desc.str.contains("Paper Bag", case=False, na=False)

        mask_rename = mask_zz & ~has_paper_bag
        if mask_rename.any():
            def get_first_word(desc: str) -> str:
                words = desc.strip().split()
                return words[0].capitalize() if words else "ZZ"
            df.loc[mask_rename, "BRAND"] = long_desc[mask_rename].apply(get_first_word)
    return df


def remove_negative_daily_stores(df: pd.DataFrame) -> pd.DataFrame:
    if not all(col in df.columns for col in ["STORE_NAME", "RCPT_DATE", "EXT_PRICE_W_TAX"]):
        return df
    daily_totals = df.groupby(["STORE_NAME", "RCPT_DATE"], dropna=False)["EXT_PRICE_W_TAX"].sum()
    negative_store_dates = daily_totals[daily_totals < 0].index
    if len(negative_store_dates) > 0:
        mask = df.set_index(["STORE_NAME", "RCPT_DATE"]).index.isin(negative_store_dates)
        df = df[~mask].reset_index(drop=True)
    return df


def remove_sale_return_pairs(df: pd.DataFrame) -> pd.DataFrame:
    if "EXT_PRICE_W_TAX" not in df.columns:
        return df
    cols_to_check = [c for c in ["RECEIPT_NO", "CUSTOMER_NAME", "GENERAL_NOTES"] if c in df.columns]
    if not cols_to_check:
        return df

    indices_to_drop = set()
    for col in cols_to_check:
        for _, group in df.groupby(col, dropna=False):
            if len(group) < 2:
                continue
            key_val = group[col].iloc[0]
            if pd.isna(key_val) or str(key_val).strip() == "":
                continue

            prices = group["EXT_PRICE_W_TAX"].values
            brands = group["BRAND"].fillna("").astype(str).str.strip().values if "BRAND" in group.columns else None
            descs = group["LONG_DESCRIPTION"].fillna("").astype(str).str.strip().values if "LONG_DESCRIPTION" in group.columns else None
            index_labels = group.index.values
            dropped_in_subgroup = set()

            for i in range(len(prices)):
                if index_labels[i] in indices_to_drop or i in dropped_in_subgroup:
                    continue
                for j in range(i + 1, len(prices)):
                    if index_labels[j] in indices_to_drop or j in dropped_in_subgroup:
                        continue
                    brand_match = (brands is None) or (brands[i] == brands[j])
                    desc_match = (descs is None) or (descs[i] == descs[j])
                    if brand_match and desc_match and np.isclose(prices[i], -prices[j], atol=1e-5) and prices[i] != 0:
                        dropped_in_subgroup.update({i, j})
                        indices_to_drop.update({index_labels[i], index_labels[j]})
                        break

    if indices_to_drop:
        df = df.drop(index=list(indices_to_drop)).reset_index(drop=True)
    return df


# =========================================================
# MAIN PROCESSING FUNCTION (Called by GUI)
# =========================================================

import pandas as pd
from pathlib import Path
import datetime
import openpyxl
import numpy as np
import re
from openpyxl.worksheet.table import Table, TableStyleInfo

# =========================================================
# CONFIGURATION & CONSTANTS
# =========================================================

BASE_FOLDER = Path(__file__).resolve().parent.parent
OUTPUT_DIR = Path.home() / "Desktop" / "Cleaned Data"

REQUIRED_COLUMNS = [
    "STORE_NAME",
    "BRAND",
    "EXT_PRICE_W_TAX",
    "RECEIPT_NO",
    "CUSTOMER_NAME",
]

COLUMNS_TO_DROP = [
    "SERIAL_NO",
    "MOBILE_NUMBER",
    "DATABASE_NAME",
    "CARD_TYPE",
    "TYPE_OF_PAYMENT",
    "TERMS",
    "AUTHORIZATION_CODE",
    "USE_CASE",
    "UNIT_ORIG_PRICE_W_TAX",
    "UNIT_DISC_AMT_W_TAX",
    "UNIT_FINAL_PRICE_W_TAX",
]

STORE_TYPO_FIXES = {
    "ALTITUDE PIONER": "ALTITUDE PIONEER",
}

DIRECT_STORES = [
    "ECOFLOW",
    "OURA",
    "RAPTOR",
    "INSTA360",
]

EF_TAG_DEFAULT = "EF Others"
I360_TAG_DEFAULT = "I360 Others"

EF_EXCLUDE_TERMS = [
    "CABLE", "CABLES", "ADAPTER", "ADAPTOR", "CHARGER", "FLASHLIGHT", 
    "BAG", "COVER", "TROLLEY", "PROTECTOR", "MOUNT", "DC"
]

BRAND_RULES = {
    "ALTITUDE PIONEER": ["DJI", "INSTA360", "ECOFLOW", "RAPTOR", "YABER", "ENABOT"],
    "ZENITH": ["RINGCONN","OURA","BOOX","COROS","PLAUDAI","GARMIN","SHOKZ","LOOP","POLAR","RAYNEO","EARQUIS"],
    "ALTITUDE": ["ECOFLOW", "INSTA360", "DJI","RINGCONN", "OURA", "BOOX", "COROS","PLAUDAI", "RAPTOR","SHOKZ","YABER","FUJIFILM", "RAYNEO","ENABOT","POLAR","GARMIN","OCCULI","EARQUIS","VULCROX","SONY"],
    "UG": ["ECOFLOW", "INSTA360", "DJI", "RINGCONN","ANKER","OURA","SHOKZ","YABER","GARMIN","PLAUDAI","RAYNEO","BOOX","COROS","POLAR","RAPTOR","OCCULLI","EARQUIS","SONY"],
}


# =========================================================
# HELPER FUNCTIONS: TEXT & HEADER PROCESSING
# =========================================================

def normalize_header(value) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip().upper().replace("\n", " ")


def clean_text(series: pd.Series) -> pd.Series:
    return (
        series.fillna("")
        .astype(str)
        .str.strip()
        .str.upper()
        .str.replace(r"\s+", " ", regex=True)
    )


def contains_term(text: str, term: str) -> bool:
    pattern = rf"(?<![A-Z0-9]){re.escape(term)}(?![A-Z0-9])"
    return bool(re.search(pattern, text))


def fix_store_typo(store: str) -> str:
    for wrong, correct in STORE_TYPO_FIXES.items():
        if store == wrong or store.startswith(wrong + " "):
            return correct + store[len(wrong):]
    return store


def read_data_file(file_path: Path, header=None, nrows=None) -> pd.DataFrame:
    if file_path.suffix.lower() in [".xls", ".xlsx"]:
        return pd.read_excel(file_path, header=header, nrows=nrows)
    return pd.read_csv(file_path, header=header, nrows=nrows)


# =========================================================
# MONOBRAND TAGGING LOGIC
# =========================================================

EF_TAG_DEFAULT = "EF Others"
I360_TAG_DEFAULT = "I360 Others"
RAPTOR_TAG_DEFAULT = "RAPTOR OTHERS"
OURA_TAG_DEFAULT = "OURA OTHERS"


def create_raptor_tag(long_description: str) -> str:
    desc = str(long_description).upper()
    if any(k in desc for k in ["TOOLBOX", "TOOL BOX"]):
        return "TOOLBOX"
    elif "COOLER" in desc:
        return "COOLERS"
    elif any(k in desc for k in ["DRINKWARE", "BOTTLE", "MUG", "TUMBLER", "FLASK", "CUP"]):
        return "DRINKWARE"
    elif any(k in desc for k in ["APPAREL", "SHIRT", "HOODIE", "HAT", "CAP", "TEE", "JACKET", "POLO", "SHORTS", "PANTS"]):
        return "APPAREL"
    elif any(k in desc for k in ["FOODWARE", "LUNCH", "CUTLERY", "PLATE", "BOWL", "SPORK", "MESS KIT"]):
        return "FOODWARE"
    elif any(k in desc for k in ["CAMPING", "OUTING", "TENT", "CHAIR", "TABLE", "LANTERN", "HAMMOCK"]):
        return "OUTING"
    elif any(k in desc for k in ["CASE", "LITE", "EXTREME", "CARRY"]):
        return "HARDWARE"
    elif any(k in desc for k in ["BAG", "BACKPACK", "POUCH", "SLING", "DUFFLE", "PACK", "CARRY", "SACK"]):
        return "BAGS"
    elif any(k in desc for k in ["STORAGE", "BOX", "CONTAINER", "BIN"]):
        return "STORAGES"
    return RAPTOR_TAG_DEFAULT


def create_ef_tag(long_description: str) -> str:
    desc = str(long_description).upper()
    accessory_terms = ["CABLE", "CABLES", "ADAPTOR", "ADAPTORS", "ADAPTER", "ADAPTERS", "FLASHLIGHT", "FLASHLIGHTS"]
    if any(term in desc for term in accessory_terms):
        return EF_TAG_DEFAULT
    if "SOLAR" in desc:
        return "Solar Panel"
    elif "DELTA" in desc:
        return "Delta Series"
    elif "RIVER" in desc:
        return "River Series"
    return EF_TAG_DEFAULT


def create_i360_tag(long_description: str) -> str:
    desc = str(long_description).upper()
    if "BUNDLE" not in desc and "LIMITED EDITION" not in desc:
        return I360_TAG_DEFAULT
    if "ACE PRO 2" in desc:
        return "Ace Pro 2"
    elif "GO ULTRA" in desc:
        return "Go Ultra"
    elif "GO 3S" in desc:
        return "Go 3s"
    elif "LUNA" in desc:
        return "Luna"
    elif "X5" in desc:
        return "X5"
    elif "X6" in desc:
        return "X6"
    return I360_TAG_DEFAULT


def create_oura_tag(long_description: str) -> str:
    desc = str(long_description).upper()
    if any(k in desc for k in ["CHARGER", "PROTECTOR", "SIZING KIT", "SIZING", "ACCESSORY", "ACCESSORIES"]):
        return "OURA ACCESSORIES"
    elif any(k in desc for k in ["RING 4", "RING4"]):
        return "OURA RING 4"
    elif any(k in desc for k in ["RING 5", "RING5"]):
        return "OURA RING 5"
    return OURA_TAG_DEFAULT


def resolve_product_tag(brand: str, long_description: str) -> str:
    bnd = str(brand).strip().upper()
    if bnd == "RAPTOR":
        return create_raptor_tag(long_description)
    elif bnd == "ECOFLOW":
        return create_ef_tag(long_description)
    elif bnd == "INSTA360":
        return create_i360_tag(long_description)
    elif bnd == "OURA":
        return create_oura_tag(long_description)
    return ""


# =========================================================
# HELPER FUNCTIONS: DATA CLEANING & FILTERING
# =========================================================

def create_channel(store: str) -> str:
    if not store:
        return ""
    if store.startswith("ALTITUDE PIONEER"):
        return "ALTITUDE PIONEER"
    return store.split()[0]


def create_category(store: str, brand: str) -> str:
    if not store:
        return ""
    for name in DIRECT_STORES:
        if contains_term(store, name):
            return name
    for store_key, allowed_brands in BRAND_RULES.items():
        if contains_term(store, store_key):
            if not brand:
                return "OTHER BRANDS"
            for item in allowed_brands:
                if contains_term(brand, item):
                    return item
            return "OTHER BRANDS"
    return ""


def clean_special_brands(df: pd.DataFrame) -> pd.DataFrame:
    if "BRAND" not in df.columns or "LONG_DESCRIPTION" not in df.columns:
        return df
    brand = df["BRAND"].fillna("").astype(str).str.strip()
    long_desc = df["LONG_DESCRIPTION"].fillna("").astype(str)

    # Reclassify JINKO entries as RAPTOR when their long description mentions Raptor.
    jinko_raptor_mask = (
        brand.str.upper().eq("JINKO")
        & long_desc.str.contains("Raptor", case=False, na=False)
    )
    df.loc[jinko_raptor_mask, "BRAND"] = "RAPTOR"

    df.loc[(brand == "*") & long_desc.str.contains("Akaso", case=False, na=False), "BRAND"] = "Akaso"
    df.loc[brand.str.upper() == "PLAUD", "BRAND"] = "PlaudAi"
    
    brand = df["BRAND"].fillna("").astype(str).str.strip()
    mask_zz = brand == "ZZ"
    if mask_zz.any():
        has_paper_bag = long_desc.str.contains("Paper Bag", case=False, na=False)
        if "EXT_PRICE_W_TAX" in df.columns:
            zero_price = df["EXT_PRICE_W_TAX"].fillna(0) == 0
            df = df[~(mask_zz & has_paper_bag & zero_price)].reset_index(drop=True)
            brand = df["BRAND"].fillna("").astype(str).str.strip()
            long_desc = df["LONG_DESCRIPTION"].fillna("").astype(str)
            mask_zz = brand == "ZZ"
            has_paper_bag = long_desc.str.contains("Paper Bag", case=False, na=False)

        mask_rename = mask_zz & ~has_paper_bag
        if mask_rename.any():
            def get_first_word(desc: str) -> str:
                words = desc.strip().split()
                return words[0].capitalize() if words else "ZZ"
            df.loc[mask_rename, "BRAND"] = long_desc[mask_rename].apply(get_first_word)
    return df


def remove_negative_daily_stores(df: pd.DataFrame) -> pd.DataFrame:
    if not all(col in df.columns for col in ["STORE_NAME", "RCPT_DATE", "EXT_PRICE_W_TAX"]):
        return df
    daily_totals = df.groupby(["STORE_NAME", "RCPT_DATE"], dropna=False)["EXT_PRICE_W_TAX"].sum()
    negative_store_dates = daily_totals[daily_totals < 0].index
    if len(negative_store_dates) > 0:
        mask = df.set_index(["STORE_NAME", "RCPT_DATE"]).index.isin(negative_store_dates)
        df = df[~mask].reset_index(drop=True)
    return df


def remove_sale_return_pairs(df: pd.DataFrame) -> pd.DataFrame:
    if "EXT_PRICE_W_TAX" not in df.columns:
        return df
    cols_to_check = [c for c in ["RECEIPT_NO", "CUSTOMER_NAME", "GENERAL_NOTES"] if c in df.columns]
    if not cols_to_check:
        return df

    indices_to_drop = set()
    for col in cols_to_check:
        for _, group in df.groupby(col, dropna=False):
            if len(group) < 2:
                continue
            key_val = group[col].iloc[0]
            if pd.isna(key_val) or str(key_val).strip() == "":
                continue

            prices = group["EXT_PRICE_W_TAX"].values
            brands = group["BRAND"].fillna("").astype(str).str.strip().values if "BRAND" in group.columns else None
            descs = group["LONG_DESCRIPTION"].fillna("").astype(str).str.strip().values if "LONG_DESCRIPTION" in group.columns else None
            index_labels = group.index.values
            dropped_in_subgroup = set()

            for i in range(len(prices)):
                if index_labels[i] in indices_to_drop or i in dropped_in_subgroup:
                    continue
                for j in range(i + 1, len(prices)):
                    if index_labels[j] in indices_to_drop or j in dropped_in_subgroup:
                        continue
                    brand_match = (brands is None) or (brands[i] == brands[j])
                    desc_match = (descs is None) or (descs[i] == descs[j])
                    if brand_match and desc_match and np.isclose(prices[i], -prices[j], atol=1e-5) and prices[i] != 0:
                        dropped_in_subgroup.update({i, j})
                        indices_to_drop.update({index_labels[i], index_labels[j]})
                        break

    if indices_to_drop:
        df = df.drop(index=list(indices_to_drop)).reset_index(drop=True)
    return df


# =========================================================
# MAIN PROCESSING FUNCTION (Called by GUI)
# =========================================================

def process_cadence(input_file_path):
    input_path = Path(input_file_path)
    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found:\n{input_path}")

    preview = read_data_file(input_path, header=None, nrows=100)
    required = set(REQUIRED_COLUMNS)
    header_row = None

    for row_index, row in preview.iterrows():
        row_values = {normalize_header(value) for value in row.tolist()}
        if required.issubset(row_values):
            header_row = row_index
            break

    if header_row is None:
        raise ValueError(f"Could not find all required columns in the first 100 rows:\n{REQUIRED_COLUMNS}")

    df = read_data_file(input_path, header=header_row)
    df.columns = df.columns.astype(str).str.strip().str.upper().str.replace(r"\s+", " ", regex=True)
    df = df.loc[:, ~df.columns.str.contains('UNNAMED', case=False)]

    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise KeyError(f"Missing required column(s): {missing}")

    df["EXT_PRICE_W_TAX"] = pd.to_numeric(df["EXT_PRICE_W_TAX"], errors="coerce").fillna(0)
    df = clean_special_brands(df)

    long_desc_clean = clean_text(df["LONG_DESCRIPTION"]) if "LONG_DESCRIPTION" in df.columns else pd.Series("", index=df.index)
    interest_mask = df["BRAND"].eq("INTEREST") | long_desc_clean.str.contains("INTEREST", case=False, na=False)

    if interest_mask.any():
        interest_df = df[interest_mask].copy()
        interest_grouped = interest_df.groupby(["RECEIPT_NO", "CUSTOMER_NAME"], as_index=False)["EXT_PRICE_W_TAX"].sum()
        interest_grouped = interest_grouped.rename(columns={"EXT_PRICE_W_TAX": "TOTAL_INTEREST"})
        df = df[~interest_mask].copy()
        df = df.merge(interest_grouped, on=["RECEIPT_NO", "CUSTOMER_NAME"], how="left")
        df["TOTAL_INTEREST"] = df["TOTAL_INTEREST"].fillna(0)
        receipt_totals = df.groupby(["RECEIPT_NO", "CUSTOMER_NAME"])["EXT_PRICE_W_TAX"].transform("sum")
        proportion = df["EXT_PRICE_W_TAX"] / receipt_totals.replace(0, 1)
        df["EXT_PRICE_W_TAX"] += df["TOTAL_INTEREST"] * proportion
        df = df.drop(columns=["TOTAL_INTEREST"])

   # Drop unneeded columns and zero-price rows (WITH INDEX RESET)
    to_drop = [c for c in COLUMNS_TO_DROP if c in df.columns]
    df = df.drop(columns=to_drop)
    df = df[df["EXT_PRICE_W_TAX"] != 0].copy().reset_index(drop=True)

    df = remove_sale_return_pairs(df)
    df = remove_negative_daily_stores(df)
    df["STORE_NAME"] = clean_text(df["STORE_NAME"]).apply(fix_store_typo)
    df["BRAND"] = clean_text(df["BRAND"])

    df["CHANNEL"] = df["STORE_NAME"].apply(create_channel)
    df["CATEGORY"] = [create_category(s, b) for s, b in zip(df["STORE_NAME"], df["BRAND"])]
    
    # Rebuild descriptions after row filtering so each brand stays aligned with its own product.
    current_long_desc = clean_text(df["LONG_DESCRIPTION"]) if "LONG_DESCRIPTION" in df.columns else pd.Series("", index=df.index)
    df["PRODUCT_TAG"] = [
        resolve_product_tag(bnd, desc)
        for bnd, desc in zip(df["BRAND"], current_long_desc)
    ]
    
    tail = ["CATEGORY", "CHANNEL", "PRODUCT_TAG"]
    df = df[[c for c in df.columns if c not in tail] + tail]

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    yesterday_date = (datetime.date.today() - datetime.timedelta(days=1)).strftime("%Y-%m-%d")
    base_filename = f"cleaned_cadence_{yesterday_date}"
    output_file = OUTPUT_DIR / f"{base_filename}.xlsx"

    counter = 1
    while output_file.exists():
        output_file = OUTPUT_DIR / f"{base_filename} ({counter}).xlsx"
        counter += 1

    df.to_excel(output_file, index=False, engine='openpyxl')

    wb = openpyxl.load_workbook(output_file)
    ws = wb.active
    ws.title = "Cleaned Cadence Data"
    
    max_row = ws.max_row
    max_column = ws.max_column
    col_letter = openpyxl.utils.get_column_letter(max_column)
    table_range = f"A1:{col_letter}{max_row}"

    tab = Table(displayName="CadenceCleanTable", ref=table_range)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", 
        showFirstColumn=False, 
        showLastColumn=False, 
        showRowStripes=True, 
        showColumnStripes=False
    )
    ws.add_table(tab)
    
    wb.save(output_file)

    return len(df), str(output_file)