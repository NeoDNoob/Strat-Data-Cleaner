import pandas as pd
from pathlib import Path
import datetime
import openpyxl
import numpy as np
import re
from openpyxl.worksheet.table import Table, TableStyleInfo

# =========================================================
# CONFIGURATION & CONSTANTS
# =========================================================

BASE_FOLDER = Path(__file__).resolve().parent.parent
OUTPUT_DIR = Path.home() / "Desktop" / "Cleaned Data"

REQUIRED_COLUMNS = [
    "STORE_NAME",
    "BRAND",
    "EXT_PRICE_W_TAX",
    "RECEIPT_NO",
    "CUSTOMER_NAME",
]

COLUMNS_TO_DROP = [
    "SERIAL_NO",
    "MOBILE_NUMBER",
    "DATABASE_NAME",
    "CARD_TYPE",
    "TYPE_OF_PAYMENT",
    "TERMS",
    "AUTHORIZATION_CODE",
    "USE_CASE",
    "UNIT_ORIG_PRICE_W_TAX",
    "UNIT_DISC_AMT_W_TAX",
    "UNIT_FINAL_PRICE_W_TAX",
]

STORE_TYPO_FIXES = {
    "ALTITUDE PIONER": "ALTITUDE PIONEER",
}

# Explicit tracked brands in canonical casing/formatting
TARGET_BRANDS = [
    "EcoFlow",
    "Insta360",
    "DJI",
    "Ringconn",
    "Oura",
    "Boox",
    "Coros",
    "PlaudAI",
    "Raptor",
    "Shokz",
    "Yaber",
    "Fujifilm",
    "Rayneo",
    "Enabot",
    "Polar",
    "Garmin",
    "Occulli",
    "Earquis",
    "Vulcrox",
    "Sony",
    "Anker",
    "Havit",
    "Loop",
]

# Case-insensitive mapping lookup for fast standard matching
BRAND_MAP = {b.upper(): b for b in TARGET_BRANDS}

# Monobrand Store Mapping Rules
MONOBRAND_KEYWORDS = {
    "ECOFLOW": "EcoFlow Stores",
    "INSTA360": "Insta360 Stores",
    "RAPTOR": "Raptor Stores",
    "OURA": "Oura Stores",
}


# =========================================================
# HELPER FUNCTIONS: TEXT & HEADER PROCESSING
# =========================================================

def normalize_header(value) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip().upper().replace("\n", " ")


def clean_text(series: pd.Series) -> pd.Series:
    return (
        series.fillna("")
        .astype(str)
        .str.strip()
        .str.upper()
        .str.replace(r"\s+", " ", regex=True)
    )


def contains_term(text: str, term: str) -> bool:
    pattern = rf"(?<![A-Z0-9]){re.escape(term)}(?![A-Z0-9])"
    return bool(re.search(pattern, text))


def fix_store_typo(store: str) -> str:
    for wrong, correct in STORE_TYPO_FIXES.items():
        if store == wrong or store.startswith(wrong + " "):
            return correct + store[len(wrong):]
    return store


def read_data_file(file_path: Path, header=None, nrows=None) -> pd.DataFrame:
    if file_path.suffix.lower() in [".xls", ".xlsx"]:
        return pd.read_excel(file_path, header=header, nrows=nrows)
    return pd.read_csv(file_path, header=header, nrows=nrows)


# =========================================================
# SEGMENTATION & QUARTER TRACKING LOGIC
# =========================================================

def segment_channel(store_name: str) -> str:
    """
    Segment stores into:
    - Altitude
    - Urban Gadgets
    - Zenith
    - Monobrand stores (EcoFlow Stores, Insta360 Stores, Raptor Stores, Oura Stores)
    """
    store = str(store_name).upper().strip()
    
    # 1. Monobrand check
    for key, label in MONOBRAND_KEYWORDS.items():
        if contains_term(store, key):
            return label
            
    # 2. Main retail brand channels
    if "ALTITUDE" in store:
        return "Altitude"
    elif "ZENITH" in store:
        return "Zenith"
    elif "URBAN GADGETS" in store or "UG " in store or store.startswith("UG") or " URBAN " in store:
        return "Urban Gadgets"
        
    return "Other Channels"


def assign_tracked_brand(brand_raw: str) -> str:
    """Categorize brand into target brand list or 'Other Brands'."""
    b_upper = str(brand_raw).strip().upper()
    return BRAND_MAP.get(b_upper, "Other Brands")


def derive_quarter(df: pd.DataFrame) -> pd.Series:
    """Extract Quarter (e.g. Q1, Q2, Q3, Q4) from RCPT_DATE if available."""
    if "RCPT_DATE" in df.columns:
        dates = pd.to_datetime(df["RCPT_DATE"], errors="coerce")
        quarters = dates.dt.quarter
        return quarters.apply(lambda q: f"Q{int(q)}" if pd.notna(q) else "Unknown Quarter")
    return pd.Series("Unknown Quarter", index=df.index)


# =========================================================
# DATA CLEANING & WORKAROUNDS
# =========================================================

def clean_special_brands(df: pd.DataFrame) -> pd.DataFrame:
    if "BRAND" not in df.columns or "LONG_DESCRIPTION" not in df.columns:
        return df
    brand = df["BRAND"].fillna("").astype(str).str.strip()
    long_desc = df["LONG_DESCRIPTION"].fillna("").astype(str)

    # Reclassify JINKO entries as RAPTOR when long description mentions Raptor.
    jinko_raptor_mask = (
        brand.str.upper().eq("JINKO")
        & long_desc.str.contains("Raptor", case=False, na=False)
    )
    df.loc[jinko_raptor_mask, "BRAND"] = "RAPTOR"

    df.loc[(brand == "*") & long_desc.str.contains("Akaso", case=False, na=False), "BRAND"] = "Akaso"
    df.loc[brand.str.upper() == "PLAUD", "BRAND"] = "PlaudAi"
    df.loc[brand.str.upper() == "PLAUDAI", "BRAND"] = "PlaudAi"
    
    brand = df["BRAND"].fillna("").astype(str).str.strip()
    mask_zz = brand == "ZZ"
    if mask_zz.any():
        has_paper_bag = long_desc.str.contains("Paper Bag", case=False, na=False)
        if "EXT_PRICE_W_TAX" in df.columns:
            zero_price = df["EXT_PRICE_W_TAX"].fillna(0) == 0
            df = df[~(mask_zz & has_paper_bag & zero_price)].reset_index(drop=True)
            brand = df["BRAND"].fillna("").astype(str).str.strip()
            long_desc = df["LONG_DESCRIPTION"].fillna("").astype(str)
            mask_zz = brand == "ZZ"
            has_paper_bag = long_desc.str.contains("Paper Bag", case=False, na=False)

        mask_rename = mask_zz & ~has_paper_bag
        if mask_rename.any():
            def get_first_word(desc: str) -> str:
                words = desc.strip().split()
                return words[0].capitalize() if words else "ZZ"
            df.loc[mask_rename, "BRAND"] = long_desc[mask_rename].apply(get_first_word)
    return df


def remove_negative_daily_stores(df: pd.DataFrame) -> pd.DataFrame:
    if not all(col in df.columns for col in ["STORE_NAME", "RCPT_DATE", "EXT_PRICE_W_TAX"]):
        return df
    daily_totals = df.groupby(["STORE_NAME", "RCPT_DATE"], dropna=False)["EXT_PRICE_W_TAX"].sum()
    negative_store_dates = daily_totals[daily_totals < 0].index
    if len(negative_store_dates) > 0:
        mask = df.set_index(["STORE_NAME", "RCPT_DATE"]).index.isin(negative_store_dates)
        df = df[~mask].reset_index(drop=True)
    return df


def remove_sale_return_pairs(df: pd.DataFrame) -> pd.DataFrame:
    if "EXT_PRICE_W_TAX" not in df.columns:
        return df
    cols_to_check = [c for c in ["RECEIPT_NO", "CUSTOMER_NAME", "GENERAL_NOTES"] if c in df.columns]
    if not cols_to_check:
        return df

    indices_to_drop = set()
    for col in cols_to_check:
        for _, group in df.groupby(col, dropna=False):
            if len(group) < 2:
                continue
            key_val = group[col].iloc[0]
            if pd.isna(key_val) or str(key_val).strip() == "":
                continue

            prices = group["EXT_PRICE_W_TAX"].values
            brands = group["BRAND"].fillna("").astype(str).str.strip().values if "BRAND" in group.columns else None
            descs = group["LONG_DESCRIPTION"].fillna("").astype(str).str.strip().values if "LONG_DESCRIPTION" in group.columns else None
            index_labels = group.index.values
            dropped_in_subgroup = set()

            for i in range(len(prices)):
                if index_labels[i] in indices_to_drop or i in dropped_in_subgroup:
                    continue
                for j in range(i + 1, len(prices)):
                    if index_labels[j] in indices_to_drop or j in dropped_in_subgroup:
                        continue
                    brand_match = (brands is None) or (brands[i] == brands[j])
                    desc_match = (descs is None) or (descs[i] == descs[j])
                    if brand_match and desc_match and np.isclose(prices[i], -prices[j], atol=1e-5) and prices[i] != 0:
                        dropped_in_subgroup.update({i, j})
                        indices_to_drop.update({index_labels[i], index_labels[j]})
                        break

    if indices_to_drop:
        df = df.drop(index=list(indices_to_drop)).reset_index(drop=True)
    return df


# =========================================================
# MAIN PROCESSING FUNCTION
# =========================================================

def process_brand_performance(input_file_path):
    input_path = Path(input_file_path)
    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found:\n{input_path}")

    preview = read_data_file(input_path, header=None, nrows=100)
    required = set(REQUIRED_COLUMNS)
    header_row = None

    for row_index, row in preview.iterrows():
        row_values = {normalize_header(value) for value in row.tolist()}
        if required.issubset(row_values):
            header_row = row_index
            break

    if header_row is None:
        raise ValueError(f"Could not find all required columns in the first 100 rows:\n{REQUIRED_COLUMNS}")

    df = read_data_file(input_path, header=header_row)
    df.columns = df.columns.astype(str).str.strip().str.upper().str.replace(r"\s+", " ", regex=True)
    df = df.loc[:, ~df.columns.str.contains('UNNAMED', case=False)]

    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise KeyError(f"Missing required column(s): {missing}")

    # Standardize numeric price field
    df["EXT_PRICE_W_TAX"] = pd.to_numeric(df["EXT_PRICE_W_TAX"], errors="coerce").fillna(0)
    
    # 1. Brand reclassification & masking fixes
    df = clean_special_brands(df)

    # 2. Interest distribution logic
    long_desc_clean = clean_text(df["LONG_DESCRIPTION"]) if "LONG_DESCRIPTION" in df.columns else pd.Series("", index=df.index)
    interest_mask = df["BRAND"].eq("INTEREST") | long_desc_clean.str.contains("INTEREST", case=False, na=False)

    if interest_mask.any():
        interest_df = df[interest_mask].copy()
        interest_grouped = interest_df.groupby(["RECEIPT_NO", "CUSTOMER_NAME"], as_index=False)["EXT_PRICE_W_TAX"].sum()
        interest_grouped = interest_grouped.rename(columns={"EXT_PRICE_W_TAX": "TOTAL_INTEREST"})
        df = df[~interest_mask].copy()
        df = df.merge(interest_grouped, on=["RECEIPT_NO", "CUSTOMER_NAME"], how="left")
        df["TOTAL_INTEREST"] = df["TOTAL_INTEREST"].fillna(0)
        receipt_totals = df.groupby(["RECEIPT_NO", "CUSTOMER_NAME"])["EXT_PRICE_W_TAX"].transform("sum")
        proportion = df["EXT_PRICE_W_TAX"] / receipt_totals.replace(0, 1)
        df["EXT_PRICE_W_TAX"] += df["TOTAL_INTEREST"] * proportion
        df = df.drop(columns=["TOTAL_INTEREST"])

    # 3. Drop unneeded columns and zero-price rows (with index reset)
    to_drop = [c for c in COLUMNS_TO_DROP if c in df.columns]
    df = df.drop(columns=to_drop)
    df = df[df["EXT_PRICE_W_TAX"] != 0].copy().reset_index(drop=True)

    # 4. Sale-return pairs offset workaround
    df = remove_sale_return_pairs(df)
    
    # 5. Negative daily store balance cleanup
    df = remove_negative_daily_stores(df)
    
    # 6. Store name normalization
    df["STORE_NAME"] = clean_text(df["STORE_NAME"]).apply(fix_store_typo)

    # 7. Apply channel segmentation, brand categorization, and quarter tracking
    df["CHANNEL_SEGMENT"] = df["STORE_NAME"].apply(segment_channel)
    df["TRACKED_BRAND"] = df["BRAND"].apply(assign_tracked_brand)
    df["QUARTER"] = derive_quarter(df)

    # Reorder key summary columns to the end
    tail = ["CHANNEL_SEGMENT", "TRACKED_BRAND", "QUARTER"]
    df = df[[c for c in df.columns if c not in tail] + tail]

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    yesterday_date = (datetime.date.today() - datetime.timedelta(days=1)).strftime("%Y-%m-%d")
    base_filename = f"cleaned_brand_performance_{yesterday_date}"
    output_file = OUTPUT_DIR / f"{base_filename}.xlsx"

    counter = 1
    while output_file.exists():
        output_file = OUTPUT_DIR / f"{base_filename} ({counter}).xlsx"
        counter += 1

    df.to_excel(output_file, index=False, engine='openpyxl')

    wb = openpyxl.load_workbook(output_file)
    ws = wb.active
    ws.title = "Brand Performance Data"
    
    max_row = ws.max_row
    max_column = ws.max_column
    col_letter = openpyxl.utils.get_column_letter(max_column)
    table_range = f"A1:{col_letter}{max_row}"

    tab = Table(displayName="BrandPerformanceTable", ref=table_range)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", 
        showFirstColumn=False, 
        showLastColumn=False, 
        showRowStripes=True, 
        showColumnStripes=False
    )
    ws.add_table(tab)
    
    wb.save(output_file)

    return len(df), str(output_file)