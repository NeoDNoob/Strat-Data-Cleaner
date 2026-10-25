from __future__ import annotations

import datetime
import re
from pathlib import Path
from typing import Any, Optional

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# Added comments so that it will push to github and trigger the workflow to build the exe file

# =========================================================
# CONFIG & MAPPINGS
# =========================================================

# Dynamic Desktop Output Directory (Works on Windows, macOS, and Linux)
OUTPUT_DIR = Path.home() / "Desktop" / "Cleaned Data"

HEADER_SCAN_ROWS = 30
MIN_HEADER_MATCHES = 5

EXPECTED_COLUMNS = [
    "DATABASE_NAME", "STORE_CODE", "STORE_NAME", "RCPT_DATE", "RECEIPT_NO",
    "CUSTOMER_NAME", "MOBILE_NUMBER", "ADDRESS_1", "TRANSACTION_TYPE", "SERIAL_NO",
    "DESCRIPTION_1", "BRAND", "LONG_DESCRIPTION", "SOLD_QTY", "TENDER_TYPE",
    "CARD_TYPE", "TYPE_OF_PAYMENT", "TERMS", "AUTHORIZATION_CODE", "USE_CASE", "TAG",
    "ASSOCIATE", "UNIT_ORIG_PRICE_W_TAX", "UNIT_DISC_AMT_W_TAX",
    "UNIT_FINAL_PRICE_W_TAX", "EXT_PRICE_W_TAX", "GENERAL_NOTES",
]

COLUMN_ALIASES = {
    "UNIT_ORGI_PRICE_W_TAX": "UNIT_ORIG_PRICE_W_TAX",
}

COLUMNS_TO_REMOVE = [
    "DATABASE_NAME", "MOBILE_NUMBER", "SERIAL_NO", "CARD_TYPE", "AUTHORIZATION_CODE",
    "USE_CASE", "UNIT_ORIG_PRICE_W_TAX", "UNIT_DISC_AMT_W_TAX", "UNIT_FINAL_PRICE_W_TAX",
    "TERMS", "TYPE_OF_PAYMENT",
]

REQUIRED_COLUMNS = ["STORE_NAME", "RCPT_DATE", "SOLD_QTY", "EXT_PRICE_W_TAX"]
TEXT_ID_COLUMNS = ["STORE_CODE", "RECEIPT_NO", "DESCRIPTION_1"]
NUMERIC_COLUMNS = ["SOLD_QTY", "EXT_PRICE_W_TAX"]

STORE_NAME_MAP = {
    "ALTITUDE ALABANG TOWN CENTER": "Altitude Alabang Town Center",
    "ALTITUDE FESTIVAL MALL": "Altitude Festival Mall",
    "ALTITUDE GATEWAY MALL": "Altitude Gateway Mall",
    "ALTITUDE GLORIETTA": "Altitude Glorietta",
    "ALTITUDE GREENBELT 5": "Altitude Greenbelt",
    "ALTITUDE MEGAMALL": "Altitude Megamall",
    "ALTITUDE PIONEER": "Altitude Pioneer",
    "ALTITUDE RPM": "Altitude Robinsons Place Manila",
    "ALTITUDE TRINOMA": "Altitude Trinoma",
    "ECOFLOW SM CITY CALAMBA": "Ecoflow SM City Calamba",
    "ECOFLOW SM CITY CLARK": "Ecoflow SM City Clark",
    "ECOFLOW SM CITY FAIRVIEW": "Ecoflow SM City Fairview",
    "ECOFLOW SM CITY TAYTAY": "Ecoflow SM City Taytay",
    "ECOFLOW SM DASMA": "Ecoflow SM Dasma",
    "ECOFLOW SM GRAND CENTRAL": "Ecoflow SM Grand Central",
    "ECOFLOW SM MALL OF ASIA": "Ecoflow SM Mall Of Asia",
    "ECOFLOW SM MARILAO": "Ecoflow SM Marilao",
    "ECOFLOW SM MEGAMALL": "Ecoflow SM Megamall",
    "ECOFLOW SM NORTH EDSA": "Ecoflow SM North Edsa",
    "ECOFLOW SM PAMPANGA": "Ecoflow SM Pampanga",
    "INSTA360 FESTIVAL MALL": "Insta360 Festival Mall",
    "INSTA360 ROBINSON MANILA": "Insta360 Robinson Manila",
    "INSTA360 SM MALL OF ASIA": "Insta360 SM Mall Of Asia",
    "INSTA360 SM MEGAMALL": "Insta360 SM Megamall",
    "INSTA360 SM NORTH EDSA": "Insta360 SM North Edsa",
    "OURA GLORIETTA 1": "Oura Glorietta",
    "OURARW": "Oura Rockwell",
    "OURA SM AURA": "Oura SM Aura",
    "OURA SM MOA": "Oura SM Mall Of Asia",
    "OURA SM NORTH EDSA": "Oura SM North Edsa",
    "RAPTOR FESTIVAL MALL": "Raptor Festival Mall",
    "RAPTOR GATEWAY": "Raptor Gateway",
    "RAPTOR GRAND CENTRAL": "Raptor Grand Central",
    "RAPTOR SM NORTH EDSA": "Raptor SM North Edsa",
    "UG ALABANG": "Urban Gadgets Alabang Town Center",
    "UG BGC": "Urban Gadgets Bonifacio Global City",
    "UG GREENHILLS": "Urban Gadgets Greenhills",
    "UG MOA": "Urban Gadgets Mall Of Asia",
    "UG VERTISNORTH": "Urban Gadgets Vertis North",
    "ZENITH EASTWOOD": "Zenith Eastwood",
    "ZENITH GLORIETTA 4": "Zenith Glorietta",
    "ZENITH PIONEER": "Zenith Pioneer",
    "ZENITH SHANGRILA": "Zenith Shangrila",
}

CHANNEL_PREFIXES = {
    "Altitude": "Altitude",
    "Insta360": "Insta360",
    "EcoFlow": "EcoFlow",
    "Oura": "Oura",
    "Raptor": "Raptor",
    "Urban Gadgets": "Urban Gadgets",
    "Zenith": "Zenith",
}

MONTH_ABBR = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MONTH_FULL = ["JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE", "JULY", "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER"]
MONTH_LOOKUP = {m.upper(): i for i, m in enumerate(MONTH_ABBR, start=1)}

ILLEGAL_XML_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
RE_DMY_TEXT = re.compile(r"^(\d{1,2})[-/. ]([A-Za-z]{3,9})[-/. ,]+(\d{2}|\d{4})$")
RE_ISO = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[T ].*)?$")
RE_SLASH = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4}|\d{2})$")
RE_SERIAL_TEXT = re.compile(r"^\d{5}(?:\.\d+)?$")


# =========================================================
# HELPERS
# =========================================================

def normalize_col_name(name: Any) -> str:
    return re.sub(r"[\s_]+", "_", str(name).strip()).strip("_").upper()

def canonical_col_name(name: Any) -> str:
    key = normalize_col_name(name)
    return COLUMN_ALIASES.get(key, key)

def normalize_store_key(name: Any) -> str:
    return " ".join(str(name).split()).upper()

def clean_cell(value: Any) -> Any:
    if isinstance(value, str):
        value = ILLEGAL_XML_CHARS.sub("", value).strip()
        return value if value != "" else None
    return value

def to_text_id(value: Any) -> Any:
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return None
    if isinstance(value, (bool, np.bool_)):
        return str(value)
    if isinstance(value, (int, np.integer)):
        return str(int(value))
    if isinstance(value, (float, np.floating)):
        return str(int(value)) if float(value).is_integer() else str(value)
    return str(value).strip()


# =========================================================
# ETL FUNCTIONS
# =========================================================

def clean_special_brands(df: pd.DataFrame) -> pd.DataFrame:
    """Applies specific brand renaming and filtering logic for 'ZZ' and 'Plaud' variants."""
    if "BRAND" not in df.columns or "LONG_DESCRIPTION" not in df.columns:
        return df

    brand = df["BRAND"].fillna("").astype(str).str.strip()
    long_desc = df["LONG_DESCRIPTION"].fillna("").astype(str)

    # 1. Normalize Plaud variants to PlaudAi
    mask_plaud = brand.str.upper().str.contains("PLAUD", na=False)
    df.loc[mask_plaud, "BRAND"] = "PlaudAi"
    
    brand = df["BRAND"].fillna("").astype(str).str.strip()

    # 2 & 3. ZZ brand rules
    mask_zz = brand.str.upper() == "ZZ"
    if mask_zz.any():
        has_paper_bag = long_desc.str.contains("Paper Bag", case=False, na=False)
        
        if "EXT_PRICE_W_TAX" in df.columns:
            zero_price = pd.to_numeric(df["EXT_PRICE_W_TAX"], errors="coerce").fillna(0) == 0
            # Delete ZZ entries containing paper bag with zero price
            df = df[~(mask_zz & has_paper_bag & zero_price)].reset_index(drop=True)
            
            # Refresh masks after dropping rows
            brand = df["BRAND"].fillna("").astype(str).str.strip()
            long_desc = df["LONG_DESCRIPTION"].fillna("").astype(str)
            mask_zz = brand.str.upper() == "ZZ"
            has_paper_bag = long_desc.str.contains("Paper Bag", case=False, na=False)

        # Rename ZZ entries that do NOT contain paper bag to the first word of long description
        mask_rename = mask_zz & ~has_paper_bag
        if mask_rename.any():
            def get_first_word(desc: str) -> str:
                words = desc.strip().split()
                return words[0].capitalize() if words else "ZZ"
            df.loc[mask_rename, "BRAND"] = long_desc[mask_rename].apply(get_first_word)

    return df


def find_header_row(raw: pd.DataFrame) -> Optional[int]:
    known = set(EXPECTED_COLUMNS) | set(COLUMN_ALIASES)
    best_idx, best_score = None, 0
    for i in range(min(HEADER_SCAN_ROWS, len(raw))):
        score = sum(1 for v in raw.iloc[i].tolist() if isinstance(v, str) and normalize_col_name(v) in known)
        if score > best_score:
            best_idx, best_score = i, score
    return best_idx if best_score >= MIN_HEADER_MATCHES else None

def prepare_sheet(raw: pd.DataFrame, sheet_name: str) -> Optional[pd.DataFrame]:
    header_idx = find_header_row(raw)
    if header_idx is None:
        return None
    names = [
        canonical_col_name(v) if isinstance(v, str) and v.strip() else f"UNNAMED_{i + 1}"
        for i, v in enumerate(raw.iloc[header_idx].tolist())
    ]
    data = raw.iloc[header_idx + 1:].copy()
    data.columns = names
    return data.reset_index(drop=True)

def parse_date_value(value: Any) -> Optional[pd.Timestamp]:
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return None
    if isinstance(value, (datetime.date, datetime.datetime, pd.Timestamp)):
        ts = pd.Timestamp(value)
        return ts.tz_localize(None).normalize() if ts.tzinfo else ts.normalize()
    if isinstance(value, (int, float, np.integer, np.floating)) and not isinstance(value, bool):
        if 1 <= float(value) < 2958466:
            return (pd.Timestamp("1899-12-30") + pd.Timedelta(days=float(value))).normalize()
        return None
    text = str(value).strip()
    if RE_SERIAL_TEXT.match(text):
        return (pd.Timestamp("1899-12-30") + pd.Timedelta(days=float(text))).normalize()
    match = RE_DMY_TEXT.match(text)
    if match:
        day, month_word, year_text = match.groups()
        word = month_word.upper()
        month = MONTH_LOOKUP.get(word[:3])
        if month is None or not MONTH_FULL[month - 1].startswith(word):
            return None
        year = int(year_text)
        year = 2000 + year if len(year_text) <= 2 else year
        try:
            return pd.Timestamp(year=year, month=month, day=int(day))
        except ValueError:
            return None
    match = RE_ISO.match(text)
    if match:
        try:
            return pd.Timestamp(year=int(match.group(1)), month=int(match.group(2)), day=int(match.group(3)))
        except ValueError:
            return None
    match = RE_SLASH.match(text)
    if match:
        first, second, year_text = int(match.group(1)), int(match.group(2)), match.group(3)
        year = int(year_text)
        year = 2000 + year if len(year_text) <= 2 else year
        try:
            return pd.Timestamp(year=year, month=second, day=first)
        except ValueError:
            return None
    return None

def remove_sale_return_pairs(df: pd.DataFrame) -> pd.DataFrame:
    """Removes offsetting sale and return pairs sharing receipt, customer, or notes, with matching brand and long_description."""
    if "EXT_PRICE_W_TAX" not in df.columns:
        return df

    cols_to_check = [c for c in ["RECEIPT_NO", "CUSTOMER_NAME", "GENERAL_NOTES"] if c in df.columns]
    if not cols_to_check:
        return df

    indices_to_drop = set()

    for col in cols_to_check:
        for _, group in df.groupby(col, dropna=False):
            if len(group) < 2:
                continue
            key_val = group[col].iloc[0]
            if pd.isna(key_val) or str(key_val).strip() == "":
                continue

            prices = group["EXT_PRICE_W_TAX"].values
            brands = group["BRAND"].fillna("").astype(str).str.strip().values if "BRAND" in group.columns else None
            descs = group["LONG_DESCRIPTION"].fillna("").astype(str).str.strip().values if "LONG_DESCRIPTION" in group.columns else None
            index_labels = group.index.values
            dropped_in_subgroup = set()

            for i in range(len(prices)):
                if index_labels[i] in indices_to_drop or i in dropped_in_subgroup:
                    continue
                for j in range(i + 1, len(prices)):
                    if index_labels[j] in indices_to_drop or j in dropped_in_subgroup:
                        continue
                    
                    brand_match = (brands is None) or (brands[i] == brands[j])
                    desc_match = (descs is None) or (descs[i] == descs[j])

                    if brand_match and desc_match and np.isclose(prices[i], -prices[j], atol=1e-5) and prices[i] != 0:
                        dropped_in_subgroup.update({i, j})
                        indices_to_drop.update({index_labels[i], index_labels[j]})
                        break

    if indices_to_drop:
        df = df.drop(index=list(indices_to_drop)).reset_index(drop=True)
    return df


# =========================================================
# MAIN PROCESSOR (Called by GUI)
# =========================================================

def process_database(input_file_path: str) -> tuple[int, str]:
    input_path = Path(input_file_path)

    # 1. Read file
    if input_path.suffix.lower() in [".xls", ".xlsx"]:
        frames = pd.read_excel(input_path, sheet_name=None, header=None, dtype=object, engine="openpyxl", keep_default_na=False, na_values=[""])
    else:
        with open(input_path, newline="", encoding="utf-8-sig", errors="replace") as f:
            rows = list(__import__('csv').reader(f))
        width = max((len(r) for r in rows), default=0)
        padded = [r + [""] * (width - len(r)) for r in rows]
        frames = {"csv": pd.DataFrame(padded, dtype=object)}

    prepared = [prepare_sheet(raw, sname) for sname, raw in frames.items()]
    prepared = [p for p in prepared if p is not None]
    if not prepared:
        raise ValueError("No valid header row found in the database export file.")

    df = pd.concat(prepared, ignore_index=True) if len(prepared) > 1 else prepared[0]

    # Clean cells and drop blank rows
    data_cols = [c for c in df.columns if not c.startswith("UNNAMED_")]
    for col in data_cols:
        df[col] = df[col].map(clean_cell)
    df = df.dropna(how="all", subset=data_cols).reset_index(drop=True)

    # Check required columns
    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise KeyError(f"Missing required column(s): {missing}")

    # Remove unwanted columns
    to_drop = [c for c in COLUMNS_TO_REMOVE if c in df.columns]
    df = df.drop(columns=to_drop)

    # Apply special brand cleaning rules for ZZ and PlaudAi
    df = clean_special_brands(df)

    # Standardize Store Names
    lookup = {normalize_store_key(k): v for k, v in STORE_NAME_MAP.items()}
    standard = {normalize_store_key(v): v for v in STORE_NAME_MAP.values()}
    original_stores = df["STORE_NAME"]
    keys = original_stores.map(lambda v: normalize_store_key(v) if isinstance(v, str) else None)
    mapped = keys.map(lookup).where(lambda s: s.notna(), keys.map(standard))
    df["STORE_NAME"] = mapped.where(mapped.notna(), original_stores)

    # Add CHANNEL column
    def channel_of(name: Any) -> Optional[str]:
        if not isinstance(name, str):
            return None
        key = normalize_store_key(name)
        return next((label for prefix, label in CHANNEL_PREFIXES.items() if key.startswith(normalize_store_key(prefix))), None)

    df.insert(0, "CHANNEL", df["STORE_NAME"].map(channel_of))

    # Parse Dates & Calendar Columns
    df["RCPT_DATE"] = pd.to_datetime([parse_date_value(v) for v in df["RCPT_DATE"]], errors="coerce")
    
    dates = df["RCPT_DATE"]
    month_no = dates.dt.month.astype("Int64")
    month_name = month_no.map(lambda m: MONTH_ABBR[int(m) - 1] if pd.notna(m) else None)
    
    # Format week number as "W" + week (e.g., W40) with safe handling for nulls
    iso_weeks = dates.dt.isocalendar().week
    week_no = pd.Series(
        [f"W{int(w)}" if pd.notna(w) else None for w in iso_weeks], index=df.index
    )
    
    quarter_year = pd.Series([f"Q{d.quarter} {d.year}" if pd.notna(d) else None for d in dates], index=df.index)

    pos = df.columns.get_loc("RCPT_DATE") + 1
    df.insert(pos, "MONTH_NAME", month_name)
    df.insert(pos + 1, "MONTH_NO", month_no)
    df.insert(pos + 2, "WEEK_NO", week_no)
    df.insert(pos + 3, "QUARTER_YEAR", quarter_year)

    # Convert Types
    for col in TEXT_ID_COLUMNS:
        if col in df.columns:
            df[col] = df[col].map(to_text_id)

    for col in NUMERIC_COLUMNS:
        if col in df.columns:
            stripped = df[col].map(lambda v: v.replace(",", "").strip() if isinstance(v, str) else v)
            df[col] = pd.to_numeric(stripped, errors="coerce")

    # --- PROPORTIONAL INTEREST ABSORPTION LOGIC ---
    interest_mask = (
        df["BRAND"].str.contains("INTEREST", case=False, na=False) |
        df["LONG_DESCRIPTION"].str.contains("INTEREST", case=False, na=False) |
        df["DESCRIPTION_1"].str.contains("INTEREST", case=False, na=False)
    )

    if interest_mask.any() and "RECEIPT_NO" in df.columns and "CUSTOMER_NAME" in df.columns:
        interest_df = df[interest_mask].copy()
        interest_grouped = interest_df.groupby(["RECEIPT_NO", "CUSTOMER_NAME"], as_index=False)["EXT_PRICE_W_TAX"].sum()
        interest_grouped = interest_grouped.rename(columns={"EXT_PRICE_W_TAX": "TOTAL_INTEREST"})

        df = df[~interest_mask].copy()

        df = df.merge(interest_grouped, on=["RECEIPT_NO", "CUSTOMER_NAME"], how="left")
        df["TOTAL_INTEREST"] = df["TOTAL_INTEREST"].fillna(0)

        receipt_totals = df.groupby(["RECEIPT_NO", "CUSTOMER_NAME"])["EXT_PRICE_W_TAX"].transform("sum")
        proportion = df["EXT_PRICE_W_TAX"] / receipt_totals.replace(0, 1)

        df["EXT_PRICE_W_TAX"] = df["EXT_PRICE_W_TAX"] + (df["TOTAL_INTEREST"] * proportion)
        df = df.drop(columns=["TOTAL_INTEREST"])

    # Remove offsetting sale-return pairs
    df = remove_sale_return_pairs(df)

    # Filter out zero price rows
    df = df[df["EXT_PRICE_W_TAX"] != 0].copy()

    # 2. Setup Destination Path with Auto-Incrementing Collision Check
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    current_date = datetime.date.today().strftime("%Y-%m-%d")
    base_filename = f"cleaned_database_{current_date}"
    output_file = OUTPUT_DIR / f"{base_filename}.xlsx"
    
    counter = 1
    while output_file.exists():
        output_file = OUTPUT_DIR / f"{base_filename} ({counter}).xlsx"
        counter += 1

    # 3. Save to Excel
    df.to_excel(output_file, index=False, engine='openpyxl')

    # 4. Format as official Excel Table
    wb = openpyxl.load_workbook(output_file)
    ws = wb.active
    
    max_row = ws.max_row
    max_col = ws.max_column
    col_letter = get_column_letter(max_col)
    table_range = f"A1:{col_letter}{max_row}"

    tab = openpyxl.worksheet.table.Table(displayName="DatabaseCleanTable", ref=table_range)
    style = openpyxl.worksheet.table.TableStyleInfo(
        name="TableStyleMedium2", 
        showFirstColumn=False, 
        showLastColumn=False, 
        showRowStripes=True, 
        showColumnStripes=False
    )
    tab.tableStyleInfo = style
    ws.add_table(tab)
    
    # Format date and currency columns
    for col_idx, col_name in enumerate(df.columns, start=1):
        col_let = get_column_letter(col_idx)
        ws.column_dimensions[col_let].width = 18
        if col_name == "RCPT_DATE":
            for cell in ws[col_let][1:]:
                cell.number_format = "mm/dd/yyyy"
        elif col_name == "EXT_PRICE_W_TAX":
            for cell in ws[col_let][1:]:
                cell.number_format = "#,##0.00"

    wb.save(output_file)

    return len(df), str(output_file)